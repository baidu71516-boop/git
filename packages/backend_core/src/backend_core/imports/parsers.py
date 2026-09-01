"""Safe tabular parsers; workbook content is treated only as untrusted data."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from backend_core.imports.enums import StoredFileType
from backend_core.imports.errors import ImportDomainError

CSV_MIME_TYPES = frozenset(
    {
        "application/csv",
        "application/octet-stream",
        "application/vnd.ms-excel",
        "text/csv",
        "text/plain",
    }
)
XLSX_MIME_TYPES = frozenset(
    {
        "application/octet-stream",
        "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
    }
)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
UNSAFE_CSV_MAGIC = (
    b"%PDF",
    b"MZ",
    b"\x7fELF",
    b"\x89PNG\r\n\x1a\n",
    b"GIF87a",
    b"GIF89a",
    b"\xff\xd8\xff",
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
    b"PK\x03\x04",
)
UNSAFE_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
UNSAFE_ARCHIVE_MARKERS = (
    "activex",
    "connections.xml",
    "customui",
    "embeddings",
    "externallinks",
    "macrosheets",
    "querytables",
    "vbaproject",
    "webextensions",
)


@dataclass(frozen=True)
class RawTabularRecord:
    row_number: int
    raw_data: dict[str, Any]
    values: dict[str, str]
    warnings: list[dict[str, str]] = field(default_factory=list)
    errors: list[dict[str, str]] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedTable:
    headers: list[str]
    rows: list[RawTabularRecord]
    file_type: StoredFileType
    mime_type: str
    encoding: str | None
    delimiter: str | None
    warnings: list[dict[str, str]]


@dataclass(frozen=True)
class ParserLimits:
    max_xlsx_uncompressed_bytes: int = 100 * 1024 * 1024
    max_xlsx_entries: int = 10_000
    max_xlsx_compression_ratio: int = 100
    max_rows: int = 100_000
    max_columns: int = 200
    max_cells: int = 5_000_000
    max_cell_chars: int = 100_000
    max_warnings: int = 10_000


def normalize_declared_mime(declared_mime: str) -> str:
    return declared_mime.partition(";")[0].strip().lower()


def _validate_declared_mime(declared_mime: str, allowed: frozenset[str]) -> None:
    normalized = normalize_declared_mime(declared_mime)
    if normalized and normalized not in allowed:
        raise ImportDomainError("MIME_MISMATCH", "Declared MIME does not match file type")


def _validate_headers(headers: list[str], limits: ParserLimits) -> None:
    if not headers or not any(item.strip() for item in headers):
        raise ImportDomainError("EMPTY_FILE", "Import file has no header")
    if len(headers) > limits.max_columns:
        raise ImportDomainError("FILE_COMPLEXITY_LIMIT", "Import file has too many columns")
    stripped = [item.strip() for item in headers]
    if any(not item for item in stripped):
        raise ImportDomainError("INVALID_HEADER", "Import header contains an empty field")
    if len(set(stripped)) != len(stripped):
        raise ImportDomainError("DUPLICATE_HEADER", "Import header contains duplicates")
    if any(len(item) > limits.max_cell_chars for item in stripped):
        raise ImportDomainError("INVALID_HEADER", "Import header exceeds the safe length")


def _validate_cell_lengths(values: list[str], limits: ParserLimits) -> list[dict[str, str]]:
    if any(len(value) > limits.max_cell_chars for value in values):
        return [{"code": "CELL_TOO_LARGE", "message": "A cell exceeds the safe length"}]
    return []


def _row_record(
    headers: list[str], values: list[str], row_number: int, limits: ParserLimits
) -> RawTabularRecord:
    padded = values[: len(headers)] + [""] * max(0, len(headers) - len(values))
    safe_values = dict(zip(headers, padded, strict=True))
    errors = _validate_cell_lengths(values, limits)
    raw_data: dict[str, Any] = dict(safe_values)
    if len(values) != len(headers):
        raw_data["__raw_values__"] = values
        errors.append(
            {
                "code": "INVALID_ROW_WIDTH",
                "message": f"Row has {len(values)} fields; expected {len(headers)}",
            }
        )
    return RawTabularRecord(
        row_number=row_number,
        raw_data=raw_data,
        values=safe_values,
        errors=errors,
    )


def _decode_csv(content: bytes) -> tuple[str, str]:
    candidate = content[3:] if content.startswith(b"\xef\xbb\xbf") else content
    if any(candidate.startswith(magic) for magic in UNSAFE_CSV_MAGIC):
        raise ImportDomainError("MIME_MISMATCH", "File content is not CSV")
    encoding = "utf-8-sig"
    try:
        text = content.decode(encoding)
    except UnicodeDecodeError:
        encoding = "gb18030"
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError as exc:
            raise ImportDomainError("INVALID_CSV_ENCODING", "Unsupported CSV encoding") from exc
    text = text.replace("\x0b", " ")
    if UNSAFE_CONTROL_CHARACTERS.search(text):
        raise ImportDomainError("INVALID_CSV", "CSV contains unsafe control characters")
    if not text.strip():
        raise ImportDomainError("EMPTY_FILE", "Import file is empty")
    return text, encoding


def _detect_delimiter(text: str) -> str:
    """Detect from the header only so a malformed data row cannot poison detection."""

    candidates = (",", ";", "\t")
    ranked: list[tuple[int, int, str]] = []
    for preference, delimiter in enumerate(candidates):
        try:
            header = next(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
        except (StopIteration, csv.Error):
            continue
        ranked.append((len(header), -preference, delimiter))
    if not ranked:
        raise ImportDomainError("INVALID_CSV", "CSV delimiter cannot be detected")
    return max(ranked)[2]


def validate_upload_type(
    content: bytes,
    *,
    suffix: str,
    declared_mime: str,
    limits: ParserLimits,
) -> StoredFileType:
    normalized_suffix = suffix.lower()
    if normalized_suffix == ".csv":
        _validate_declared_mime(declared_mime, CSV_MIME_TYPES)
        _decode_csv(content)
        return StoredFileType.CSV
    if normalized_suffix == ".xlsx":
        _validate_declared_mime(declared_mime, XLSX_MIME_TYPES)
        if not content.startswith(b"PK\x03\x04"):
            raise ImportDomainError("MIME_MISMATCH", "File content is not XLSX")
        _validate_xlsx_archive(content, limits)
        return StoredFileType.XLSX
    raise ImportDomainError("INVALID_FILE_EXTENSION", "Only CSV and XLSX are allowed")


def parse_csv(
    content: bytes, declared_mime: str, limits: ParserLimits | None = None
) -> ParsedTable:
    active_limits = limits or ParserLimits()
    _validate_declared_mime(declared_mime, CSV_MIME_TYPES)
    text, encoding = _decode_csv(content)
    try:
        delimiter = _detect_delimiter(text)
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        first = next(reader, None)
        if first is None:
            raise ImportDomainError("EMPTY_FILE", "Import file is empty")
        headers = [item.strip() for item in first]
        _validate_headers(headers, active_limits)
        rows: list[RawTabularRecord] = []
        cell_count = len(headers)
        for row_number, values in enumerate(reader, start=2):
            if len(rows) >= active_limits.max_rows:
                raise ImportDomainError("FILE_COMPLEXITY_LIMIT", "Import file has too many rows")
            cell_count += len(values)
            if cell_count > active_limits.max_cells:
                raise ImportDomainError("FILE_COMPLEXITY_LIMIT", "Import file has too many cells")
            rows.append(_row_record(headers, values, row_number, active_limits))
    except ImportDomainError:
        raise
    except csv.Error as exc:
        raise ImportDomainError("INVALID_CSV", "CSV structure is invalid") from exc
    if not rows:
        raise ImportDomainError("NO_DATA_ROWS", "Import file has no data rows")
    return ParsedTable(
        headers=headers,
        rows=rows,
        file_type=StoredFileType.CSV,
        mime_type="text/csv",
        encoding=encoding,
        delimiter=delimiter,
        warnings=[],
    )


def _validate_xlsx_archive(content: bytes, limits: ParserLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_xlsx_entries:
                raise ImportDomainError("UNSAFE_XLSX", "Workbook contains too many entries")
            seen_paths: set[str] = set()
            total = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ImportDomainError("UNSAFE_XLSX", "Encrypted workbooks are not allowed")
                if "\\" in entry.filename:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook contains an unsafe path")
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook contains an unsafe path")
                normalized_path = str(path).lower()
                if normalized_path in seen_paths:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook contains duplicate entries")
                seen_paths.add(normalized_path)
                unix_mode = entry.external_attr >> 16
                if unix_mode & 0o170000 == 0o120000:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook contains a symbolic link")
                if any(marker in normalized_path for marker in UNSAFE_ARCHIVE_MARKERS):
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook contains active content")
                total += entry.file_size
                if total > limits.max_xlsx_uncompressed_bytes:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook expands beyond the safe limit")
                if entry.file_size and not entry.compress_size:
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook has an unsafe archive entry")
                if (
                    entry.compress_size
                    and entry.file_size / entry.compress_size > limits.max_xlsx_compression_ratio
                ):
                    raise ImportDomainError("UNSAFE_XLSX", "Workbook compression ratio is unsafe")
                if normalized_path.endswith((".xml", ".rels")):
                    xml_content = archive.read(entry)
                    lowered = xml_content.lower()
                    if b"<!doctype" in lowered or b"<!entity" in lowered:
                        raise ImportDomainError("UNSAFE_XLSX", "Workbook contains unsafe XML")
                    if normalized_path.endswith(".rels") and b'targetmode="external"' in lowered:
                        raise ImportDomainError("UNSAFE_XLSX", "Workbook contains external links")
            required = {"[content_types].xml", "xl/workbook.xml"}
            if not required.issubset(seen_paths):
                raise ImportDomainError("INVALID_XLSX", "Workbook is missing required content")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"vbaproject" in content_types:
                raise ImportDomainError("UNSAFE_XLSX", "Macro-enabled workbooks are not allowed")
    except ImportDomainError:
        raise
    except (KeyError, zipfile.BadZipFile) as exc:
        raise ImportDomainError("INVALID_XLSX", "Workbook structure is invalid") from exc


def _huitun_dimension_repair_needed(content: bytes, workbook: Workbook, worksheet: Any) -> bool:
    """Recognize only Huitun's one-sheet stale ``A1`` dimension metadata."""

    if len(workbook.worksheets) != 1 or worksheet.calculate_dimension() not in {"A1", "A1:A1"}:
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    except (KeyError, ValueError, zipfile.BadZipFile):
        return False
    return re.search(rb'<c\b[^>]*\br="(?!A1")[A-Z]+[1-9][0-9]*"', sheet_xml) is not None


def parse_xlsx(
    content: bytes,
    declared_mime: str,
    limits: ParserLimits | None = None,
    *,
    repair_huitun_dimensions: bool = False,
) -> ParsedTable:
    active_limits = limits or ParserLimits()
    _validate_declared_mime(declared_mime, XLSX_MIME_TYPES)
    if not content.startswith(b"PK\x03\x04"):
        raise ImportDomainError("MIME_MISMATCH", "File content is not XLSX")
    _validate_xlsx_archive(content, active_limits)
    workbook: Workbook | None = None
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        worksheet = workbook.active
        if worksheet is None:
            raise ImportDomainError("EMPTY_FILE", "Workbook has no active worksheet")
        if repair_huitun_dimensions and _huitun_dimension_repair_needed(
            content, workbook, worksheet
        ):
            worksheet.reset_dimensions()
        iterator = worksheet.iter_rows()
        first = next(iterator, None)
        if first is None:
            raise ImportDomainError("EMPTY_FILE", "Workbook is empty")
        if any(cell.data_type == "f" for cell in first):
            raise ImportDomainError("INVALID_HEADER", "Formula cells are not allowed in headers")
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in first]
        _validate_headers(headers, active_limits)
        rows: list[RawTabularRecord] = []
        warning_count = 0
        cell_count = len(headers)
        warnings_truncated = False
        for row_number, cells in enumerate(iterator, start=2):
            if len(rows) >= active_limits.max_rows:
                raise ImportDomainError("FILE_COMPLEXITY_LIMIT", "Import file has too many rows")
            cell_count += len(cells)
            if cell_count > active_limits.max_cells:
                raise ImportDomainError("FILE_COMPLEXITY_LIMIT", "Import file has too many cells")
            raw_values: list[str] = []
            safe_values: list[str] = []
            row_warnings: list[dict[str, str]] = []
            for column, cell in enumerate(cells, start=1):
                raw_value = "" if cell.value is None else str(cell.value)
                raw_values.append(raw_value)
                if cell.data_type == "f":
                    safe_values.append("")
                    if warning_count < active_limits.max_warnings:
                        row_warnings.append(
                            {"code": "FORMULA_IGNORED", "location": f"{row_number}:{column}"}
                        )
                        warning_count += 1
                    else:
                        warnings_truncated = True
                else:
                    safe_values.append(raw_value)
            record = _row_record(headers, safe_values, row_number, active_limits)
            raw_data = dict(record.raw_data)
            for index, header in enumerate(headers):
                if index < len(raw_values):
                    raw_data[header] = raw_values[index]
            if len(raw_values) != len(headers):
                raw_data["__raw_values__"] = raw_values
            rows.append(
                RawTabularRecord(
                    row_number=row_number,
                    raw_data=raw_data,
                    values=record.values,
                    warnings=row_warnings,
                    errors=record.errors,
                )
            )
    except ImportDomainError:
        raise
    except Exception as exc:
        raise ImportDomainError("INVALID_XLSX", "Workbook cannot be parsed safely") from exc
    finally:
        if workbook is not None:
            workbook.close()
    if not rows:
        raise ImportDomainError("NO_DATA_ROWS", "Import file has no data rows")
    table_warnings = (
        [{"code": "WARNINGS_TRUNCATED", "message": "Additional warnings were suppressed"}]
        if warnings_truncated
        else []
    )
    return ParsedTable(
        headers=headers,
        rows=rows,
        file_type=StoredFileType.XLSX,
        mime_type=XLSX_MIME,
        encoding=None,
        delimiter=None,
        warnings=table_warnings,
    )


def parse_table(
    content: bytes,
    *,
    file_type: StoredFileType,
    declared_mime: str,
    limits: ParserLimits,
    repair_huitun_dimensions: bool = False,
) -> ParsedTable:
    if file_type == StoredFileType.CSV:
        return parse_csv(content, declared_mime, limits)
    if file_type == StoredFileType.XLSX:
        return parse_xlsx(
            content,
            declared_mime,
            limits,
            repair_huitun_dimensions=repair_huitun_dimensions,
        )
    raise ImportDomainError("INVALID_FILE_TYPE", "Unsupported import file type")
